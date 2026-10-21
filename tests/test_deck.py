"""Tests for the generated ForgeFlow first-version deck."""

from __future__ import annotations

import hashlib
import re
import subprocess
import sys
from pathlib import Path

import pytest

# python-pptx is a tool dependency for checking the generated deck, not a
# runtime dependency of the product. A clone without it must still be able to
# collect and run the rest of the suite.
pptx = pytest.importorskip("pptx", reason="python-pptx is needed to inspect the generated deck")
from pptx.enum.shapes import MSO_SHAPE_TYPE  # noqa: E402 - after the import guard
Presentation = pptx.Presentation

REPO = Path(__file__).resolve().parents[1]
BUILDER = REPO / "scripts" / "build_deck.py"
COMMITTED_DECK = REPO / "docs" / "ForgeFlow-deck-18-pages.pptx"
SLIDE_COUNT = 18
HEX_40 = re.compile(r"[0-9a-fA-F]{40,}")
NUMERIC_SLIDES = (3, 6, 7, 8, 10, 11, 12, 14, 15, 16)


def build_deck(out: Path) -> Path:
    result = subprocess.run(
        [sys.executable, str(BUILDER), "--out", str(out)],
        cwd=REPO,
        capture_output=True,
        text=True,
        check=False,
    )
    assert result.returncode == 0, result.stderr
    assert out.is_file()
    return out


def slide_texts(deck: Path) -> list[str]:
    presentation = Presentation(deck)
    texts = []
    for slide in presentation.slides:
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text)
            if shape.has_table:
                for row in shape.table.rows:
                    parts.extend(cell.text for cell in row.cells)
        texts.append("\n".join(parts))
    return texts


@pytest.fixture(scope="module")
def built_deck(tmp_path_factory) -> Path:
    return build_deck(tmp_path_factory.mktemp("deck") / "built.pptx")


def test_builder_writes_exactly_eighteen_slides(built_deck):
    assert len(Presentation(built_deck).slides) == SLIDE_COUNT


def test_cover_contains_project_name(built_deck):
    assert "ForgeFlow" in slide_texts(built_deck)[0]


def test_result_and_permission_numbers_are_present(built_deck):
    texts = slide_texts(built_deck)
    assert any("100.0" in text for text in texts)
    assert any("403" in text for text in texts)


def test_architecture_slide_uses_drawn_shapes_not_a_picture(built_deck):
    slide = Presentation(built_deck).slides[4]
    assert len(slide.shapes) >= 6
    assert all(shape.shape_type != MSO_SHAPE_TYPE.PICTURE for shape in slide.shapes)


def test_no_secret_like_or_host_path_text(built_deck):
    joined = "\n".join(slide_texts(built_deck))
    assert "sk-" not in joined
    assert HEX_40.search(joined) is None
    assert "/Users/" not in joined
    assert "/home/" not in joined


def test_numeric_slides_cite_evidence_and_configuration(built_deck):
    texts = slide_texts(built_deck)
    for page in NUMERIC_SLIDES:
        text = texts[page - 1]
        assert "证据" in text or "依据" in text, f"slide {page} has no evidence line"
        assert "配置" in text, f"slide {page} has no configuration line"
        assert "planning-metadata-v1" in text or "permission-split-v1" in text, (
            f"slide {page} does not name a recorded configuration"
        )


def test_committed_deck_opens_with_eighteen_slides():
    assert COMMITTED_DECK.is_file(), "run `uv run python scripts/build_deck.py` first"
    assert len(Presentation(COMMITTED_DECK).slides) == SLIDE_COUNT
    assert COMMITTED_DECK.stat().st_size < 50 * 1024 * 1024


def test_deck_is_sixteen_by_nine(built_deck):
    presentation = Presentation(built_deck)
    ratio = presentation.slide_width / presentation.slide_height
    assert abs(ratio - 16 / 9) < 0.01


def test_output_is_byte_deterministic(tmp_path):
    first = build_deck(tmp_path / "first.pptx")
    second = build_deck(tmp_path / "second.pptx")
    first_digest = hashlib.sha256(first.read_bytes()).hexdigest()
    second_digest = hashlib.sha256(second.read_bytes()).hexdigest()
    assert first_digest == second_digest
