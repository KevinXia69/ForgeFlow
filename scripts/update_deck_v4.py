#!/usr/bin/env python3
"""Update the first-release deck with the v4 evidence; keep the old deck intact."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "ForgeFlow_首版_18页.pptx"
TARGET = ROOT / "docs" / "ForgeFlow_v4_更新版.pptx"

REPLACEMENTS = {
    "保留任务 0/1": "G4 新保留 2/4 strict",
    "开发场景 A 3/3、B 2/3；保留任务 2281 一次 0 分，被原型边界拒绝。":
        "开发 A 3/3、B 2/3；2281 修复后回归 100；G4 新保留 2/4 strict。",
    "首版控制台已含运行、计划、审批与项目问答；assistant-ui 多轮对话仍在计划中":
        "工作台已含运行、计划、审批、证据面板与多轮项目问答；平台模型与自有 key 均可配置。",
    "部署与安全：本地只读已闭环，公网入口尚未完成":
        "部署与安全：本地执行闭环，公网只读评审已上线",
    "本地只读闭环": "本地执行闭环",
    "公网未部署。": "公网只读评审已部署（forgeflow-yzo9.onrender.com）。",
    "公网与完整工作台未完成。": "公网评审与完整工作台已上线。",
    "所有结果绑定 run_id；未测量项保持未知。A 3/3、B 2/3；保留任务一次 0 分。":
        "所有结果绑定 run_id；未测量项保持未知。A/B 回归 100；G4 2/4 strict；2288 超时保留。",
}


def replace_text(shape, old: str, new: str) -> bool:
    if not shape.has_text_frame:
        return False
    changed = False
    for paragraph in shape.text_frame.paragraphs:
        if old in paragraph.text:
            # Collapse to a single run so the replacement is legible even when
            # PowerPoint split the original sentence across runs.
            first = paragraph.runs[0] if paragraph.runs else None
            text = paragraph.text.replace(old, new)
            for run in list(paragraph.runs)[1:]:
                run._r.getparent().remove(run._r)
            if first is None:
                paragraph.text = text
            else:
                first.text = text
            changed = True
    return changed


def add_summary(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.7), Inches(0.6), Inches(11.9), Inches(6.6))
    frame = box.text_frame
    frame.word_wrap = True
    title = frame.paragraphs[0]
    title.text = "v4 更新：有边界的泛化与产品闭环"
    title.font.size = Pt(28)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0x1A, 0x1A, 0x18)
    lines = [
        "",
        "G0  实验清单、配置指纹、双口径与暴露登记：每次尝试可追到配置，超时与失败保留。",
        "G1  2281 归因：MOQ 合法剩余 + 组件供应商 offer 解析；修复后 2281/A/B 均 100、strict。",
        "G2  2288 同预算对照：0 次 prepare、0 写入；失败提前到上下文/输出上限，未把 0 分说成已解决。",
        "G3  逐任务合同：2281/2287 已写入审阅合同，执行器启动带 policy，plan 记录 policy_sha256。",
        "G4  新保留：预登记 4 道；2292、2290 满分 strict；2283 98.5（1 条 origin）；2285 0（scope）。",
        "D2  工作台：多轮上下文（最多 6 轮）+ 证据面板；平台默认 Nemotron，可切自有 DeepSeek key。",
        "R   结论：同基准同族有初步泛化，2/4 新保留 strict；不外推全部 repair_plan。",
    ]
    for line in lines:
        p = frame.add_paragraph()
        p.text = line
        p.font.size = Pt(15)
        p.space_after = Pt(8)


def main() -> None:
    prs = Presentation(SOURCE)
    hits = 0
    for slide in prs.slides:
        for shape in slide.shapes:
            for old, new in REPLACEMENTS.items():
                if replace_text(shape, old, new):
                    hits += 1
    add_summary(prs)
    prs.save(TARGET)
    print(f"updated {hits} text blocks -> {TARGET}")


if __name__ == "__main__":
    main()
